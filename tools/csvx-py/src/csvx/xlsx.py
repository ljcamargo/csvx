"""XLSX adapters. They never evaluate formulas."""
from __future__ import annotations

from datetime import date, datetime, time
from collections import Counter
from pathlib import Path
from typing import Literal
from zipfile import BadZipFile, is_zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.comments import Comment

from .model import Cell, ConversionReport, Document, Sheet
from .parser import parse, render


class ConversionError(RuntimeError):
    pass


def _issue(report: ConversionReport, strict: bool, feature: str, message: str, sheet: str,
           cell: str | None = None, action: str = "dropped") -> None:
    report.add(feature, message, sheet=sheet, cell=cell, action=action)
    if strict:
        raise ConversionError(f"{sheet}{'!' + cell if cell else ''}: {message}")


def _dtype(value: object) -> str | None:
    if isinstance(value, bool): return "bool"
    if isinstance(value, int): return "int"
    if isinstance(value, float): return "float"
    if isinstance(value, datetime): return "datetime"
    if isinstance(value, date): return "date"
    if isinstance(value, time): return "time"
    return None


def read_xlsx(path: str | Path, *, strict: bool = False,
              dimension: Literal["content", "worksheet", "compact", "fidelity"] = "content") -> tuple[Document, ConversionReport]:
    """Read exposed XLSX values/formulas/comments without formula evaluation."""
    path = Path(path)
    if path.suffix.lower() not in {".xlsx", ".xlsm"}: raise ConversionError("only .xlsx is supported")
    if path.suffix.lower() == ".xlsm": raise ConversionError("macro-enabled workbooks are unsupported")
    if not is_zipfile(path):
        raise ConversionError(f"{path.name} has an .xlsx extension but is not an XLSX ZIP archive; it may be CSV text")
    try:
        formulas = load_workbook(path, data_only=False, read_only=False)
        cached = load_workbook(path, data_only=True, read_only=False)
    except (BadZipFile, InvalidFileException) as error:
        raise ConversionError(f"cannot read XLSX workbook: {error}") from error
    report, sheets = ConversionReport(), []
    declarations = []
    for ws, cached_ws in zip(formulas.worksheets, cached.worksheets):
        if ws.merged_cells.ranges: _issue(report, strict, "merged_cells", "merged cells are not yet converted", ws.title)
        if ws.conditional_formatting: _issue(report, strict, "conditional_formatting", "conditional formatting is not yet converted", ws.title)
        if ws.data_validations.dataValidation: _issue(report, strict, "data_validation", "data validation is not yet converted", ws.title)
        max_row, max_col = ws.max_row, ws.max_column
        # Content mode keeps A1-relative positions and interior gaps, but
        # trims a style-only tail. The legacy compact/fidelity names remain
        # accepted during this pre-alpha transition.
        if dimension in {"content", "compact"}:
            occupied = [(c.row, c.column) for row in ws.iter_rows() for c in row
                        if c.value is not None or c.comment or c.hyperlink]
            max_row = max((r for r, _ in occupied), default=1)
            max_col = max((c for _, c in occupied), default=1)
        table = {"header": True}
        sheet = Sheet(ws.title, table=table)
        style_counts: Counter[int] = Counter()
        for row_index in range(1, max_row + 1):
            row: list[Cell] = []
            for col_index in range(1, max_col + 1):
                cell, cached_cell = ws.cell(row_index, col_index), cached_ws.cell(row_index, col_index)
                formula = cell.value if cell.data_type == "f" else None
                value = cached_cell.value if formula else cell.value
                if formula and value is not None: value = str(value)
                elif value is not None and not isinstance(value, str): value = str(value)
                output = Cell(value=value, formula=formula, dtype=_dtype(cell.value if not formula else cached_cell.value))
                if cell.comment: output.comment = cell.comment.text
                if cell.hyperlink: output.annotation = f"link: '{cell.hyperlink.target}'"
                if cell.style_id and cell.style_id != 0:
                    style_counts[cell.style_id] += 1
                row.append(output)
            if row_index == 1: sheet.header = row
            else: sheet.rows.append(row)
        if style_counts:
            _issue(report, strict, "cell_style", f"{sum(style_counts.values())} styled cells across {len(style_counts)} style records are not yet converted", ws.title)
        sheets.append(sheet)
        declarations.append({"name": ws.title})
        report.sheets_converted += 1
    # A non-default single-sheet name is workbook data and needs frontmatter;
    # `Sheet1` is the CSVX implicit default and can remain body-only.
    frontmatter = {"sheets": declarations} if len(sheets) > 1 or sheets[0].name != "Sheet1" else {}
    return Document(frontmatter=frontmatter, sheets=sheets), report


def _xlsx_value(source: Cell, report: ConversionReport, strict: bool, sheet: str, coordinate: str) -> object:
    """Coerce only unambiguous CSVX registry values for XLSX output."""
    if source.value is None:
        return None
    dtype = (source.dtype or "").split("(", 1)[0]
    try:
        if dtype == "int": return int(source.value)
        if dtype == "float": return float(source.value)
        if dtype == "bool": return source.value.lower() in {"true", "1", "yes", "on"}
        if dtype == "date": return date.fromisoformat(source.value)
        if dtype == "datetime": return datetime.fromisoformat(source.value)
        if dtype == "time": return time.fromisoformat(source.value)
    except ValueError:
        _issue(report, strict, "typed_value", f"cannot coerce {dtype} value; writing text", sheet, coordinate, action="changed")
    return source.value


def write_xlsx(document: Document, path: str | Path, *, strict: bool = False,
               materialize_dimensions: bool = False) -> ConversionReport:
    """Write values, formula text, comments, and basic date/number types to XLSX."""
    report, workbook = ConversionReport(), Workbook()
    workbook.remove(workbook.active)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    for sheet in document.sheets:
        ws = workbook.create_sheet(sheet.name)
        cached_formula_count = 0
        rows = ([sheet.header] if sheet.header is not None else []) + sheet.rows
        for r, row in enumerate(rows, 1):
            for c, source in enumerate(row, 1):
                target = ws.cell(r, c)
                if source.formula is not None:
                    if not source.formula.startswith("="):
                        _issue(report, strict, "formula", "formula does not begin with '='", sheet.name, target.coordinate)
                    target.value = source.formula
                    if source.value is not None:
                        cached_formula_count += 1
                else:
                    target.value = _xlsx_value(source, report, strict, sheet.name, target.coordinate)
                if source.comment: target.comment = Comment(source.comment, "CSVX")
                if source.annotation:
                    try:
                        annotation = __import__("yaml").safe_load(source.annotation)
                        if isinstance(annotation, dict) and isinstance(annotation.get("link") or annotation.get("url"), str):
                            target.hyperlink = annotation.get("link") or annotation.get("url")
                        else: _issue(report, strict, "annotation", "annotation is not mapped to XLSX", sheet.name, target.coordinate)
                    except Exception:
                        _issue(report, strict, "annotation", "annotation is not valid YAML", sheet.name, target.coordinate)
                if source.styles: _issue(report, strict, "cell_style", "styles are not yet converted", sheet.name, target.coordinate)
        if cached_formula_count:
            _issue(report, strict, "formula_cache", f"{cached_formula_count} cached formula values are not written; workbook will recalculate", sheet.name, action="changed")
        if materialize_dimensions:
            rows_n, cols_n = sheet.table.get("rows"), sheet.table.get("cols")
            offset = 1 if sheet.header is not None else 0
            if isinstance(rows_n, int) and isinstance(cols_n, int): ws.cell(rows_n + offset, cols_n)
        report.sheets_converted += 1
    workbook.save(path)
    return report


def xlsx_to_csvx(input_path: str | Path, output_path: str | Path, **options: object) -> ConversionReport:
    document, report = read_xlsx(input_path, strict=bool(options.get("strict", False)), dimension=options.get("dimension", "content"))
    Path(output_path).write_text(render(document), encoding="utf-8")
    return report


def csvx_to_xlsx(input_path: str | Path, output_path: str | Path, **options: object) -> ConversionReport:
    document = parse(Path(input_path).read_text(encoding="utf-8"), strict=bool(options.get("strict", False)))
    return write_xlsx(document, output_path, strict=bool(options.get("strict", False)), materialize_dimensions=bool(options.get("materialize_dimensions", False)))
