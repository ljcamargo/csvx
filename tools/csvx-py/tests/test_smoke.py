from pathlib import Path
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from csvx import ConversionError, parse, read_xlsx, xlsx_to_csvx, csvx_to_xlsx


def test_flat_frontmatter_and_formula_protection():
    doc = parse('--- {delimiter: "|"} ---\na|b\nx|"`=CONCAT(\"a|b\",\"c\")`"\n')
    assert doc.sheets[0].rows[0][1].formula == '=CONCAT("a|b","c")'


def test_xlsx_round_trip(tmp_path: Path):
    source = tmp_path / "source.xlsx"
    workbook = Workbook(); sheet = workbook.active; sheet.title = "Data"
    sheet.append(["name", "value"]); sheet.append(["Ada", 42]); sheet["B2"] = "=40+2"
    workbook.save(source)
    csvx_path, output = tmp_path / "out.csvx", tmp_path / "out.xlsx"
    xlsx_to_csvx(source, csvx_path)
    csvx_to_xlsx(csvx_path, output)
    rebuilt = load_workbook(output, data_only=False)
    assert rebuilt.active.title == "Data"
    assert rebuilt.active["B2"].value == "=40+2"


def test_xlsx_content_dimension_trims_style_only_tail_and_aggregates_styles(tmp_path: Path):
    source = tmp_path / "styled.xlsx"
    workbook = Workbook(); sheet = workbook.active
    sheet.append(["name"]); sheet.append(["Ada"])
    sheet["A2"].font = Font(bold=True)
    sheet["A100"].number_format = "0.00"  # Extends XLSX's reported range only.
    workbook.save(source)

    document, report = read_xlsx(source)

    assert len(document.sheets[0].rows) == 1
    assert len(report.warnings) == 1
    assert report.warnings[0].feature == "cell_style"
    assert "1 styled cells" in report.warnings[0].message


def test_false_xlsx_extension_has_a_clean_error(tmp_path: Path):
    source = tmp_path / "not-a-workbook.xlsx"
    source.write_text("name,value\nAda,42\n", encoding="utf-8")
    with pytest.raises(ConversionError, match="not an XLSX ZIP archive"):
        read_xlsx(source)
